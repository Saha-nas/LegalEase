import streamlit as st
import requests
from io import BytesIO
from docx import Document
from fpdf import FPDF


# --------------------------------------------------
# Page Configuration
# --------------------------------------------------

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide"
)


# --------------------------------------------------
# Title
# --------------------------------------------------

st.title("⚖️ LegalEase")
st.subheader("AI-Powered Legal Document Generator")

st.write(
    "Generate customizable legal documents using artificial intelligence."
)


# --------------------------------------------------
# Document Type
# --------------------------------------------------

document_type = st.selectbox(
    "Document Type",
    [
        "Non-Disclosure Agreement",
        "Employment Contract",
        "Freelance Work Contract",
        "Residential Lease Agreement",
        "Other"
    ]
)


# --------------------------------------------------
# Parties
# --------------------------------------------------

parties = st.text_area(
    "Parties",
    placeholder="Example: Bhuvan (Freelancer), ABC Company (Client)"
)


# --------------------------------------------------
# Terms and Conditions
# --------------------------------------------------

terms = st.text_area(
    "Terms and Conditions",
    placeholder="Enter the important terms and conditions..."
)


# --------------------------------------------------
# Date
# --------------------------------------------------

dates = st.text_input(
    "Date",
    placeholder="Example: September 23, 2026"
)


# --------------------------------------------------
# Generate Document
# --------------------------------------------------

if st.button("🚀 Generate Document", type="primary"):

    if not parties or not terms or not dates:

        st.warning("Please fill in all the fields.")

    else:

        with st.spinner("Generating your legal document..."):

            try:

                response = requests.post(
                    "https://legalease-yg44.onrender.com/generate",
                    json={
                        "document_type": document_type,
                        "parties": parties,
                        "terms": terms,
                        "dates": dates
                    }
                )


                # --------------------------------------------------
                # Successful Response
                # --------------------------------------------------

                if response.status_code == 200:

                    result = response.json()

                    document = result["document"]

                    st.success(
                        "Document generated successfully! 🎉"
                    )

                    st.subheader("Generated Document")


                    # --------------------------------------------------
                    # Preview / Edit Document
                    # --------------------------------------------------

                    edited_document = st.text_area(
                        "Preview / Edit Document",
                        value=document,
                        height=500
                    )


                    # --------------------------------------------------
                    # Download TXT
                    # --------------------------------------------------

                    st.download_button(
                        label="📄 Download TXT",
                        data=edited_document,
                        file_name="legal_document.txt",
                        mime="text/plain",
                        on_click="ignore"
                    )


                    # --------------------------------------------------
                    # Create DOCX
                    # --------------------------------------------------

                    doc = Document()

                    for paragraph in edited_document.split("\n"):
                        doc.add_paragraph(paragraph)

                    docx_file = BytesIO()

                    doc.save(docx_file)

                    docx_file.seek(0)


                    # --------------------------------------------------
                    # Download DOCX
                    # --------------------------------------------------

                    st.download_button(
                        label="📝 Download DOCX",
                        data=docx_file,
                        file_name="legal_document.docx",
                        mime=(
                            "application/vnd.openxmlformats-"
                            "officedocument.wordprocessingml.document"
                        ),
                        on_click="ignore"
                    )


                    # --------------------------------------------------
                    # Create PDF
                    # --------------------------------------------------

                    pdf = FPDF()

                    pdf.set_auto_page_break(
                        auto=True,
                        margin=15
                    )

                    pdf.add_page()


                    # --------------------------------------------------
                    # PDF Title
                    # --------------------------------------------------

                    pdf.set_font(
                        "Helvetica",
                        "B",
                        16
                    )

                    pdf.cell(
                        0,
                        10,
                        "LegalEase - Legal Document",
                        ln=1,
                        align="C"
                    )

                    pdf.ln(5)


                    # --------------------------------------------------
                    # PDF Body
                    # --------------------------------------------------

                    pdf.set_font(
                        "Helvetica",
                        "",
                        11
                    )

                    for line in edited_document.split("\n"):

                        # Remove simple Markdown formatting
                        line = line.replace("**", "")
                        line = line.replace("### ", "")
                        line = line.replace("## ", "")
                        line = line.replace("# ", "")

                        # Convert unsupported characters
                        line = line.encode(
                            "latin-1",
                            "replace"
                        ).decode("latin-1")

                        if line.strip():

                           pdf.multi_cell(w=pdf.epw, h=7, text=line, wrapmode="CHAR")

                        else:

                            pdf.ln(4)


                    # --------------------------------------------------
                    # Convert PDF to bytes
                    # --------------------------------------------------

                    pdf_output = pdf.output(dest='S')

                    if isinstance(pdf_output,str):
                        pdf_bytes = pdf_output.encode("latin-1")
                    else:
                        pdf_bytes = bytes(pdf_output)

                    


                    # --------------------------------------------------
                    # Download PDF
                    # --------------------------------------------------

                    st.download_button(
                        label="📕 Download PDF",
                        data=pdf_bytes,
                        file_name="legal_document.pdf",
                        mime="application/pdf",
                        on_click="ignore"
                    )


                # --------------------------------------------------
                # Backend Error
                # --------------------------------------------------

                else:

                    st.error(
                        f"Backend error: {response.status_code}"
                    )


            # --------------------------------------------------
            # Connection Error
            # --------------------------------------------------

            except requests.exceptions.ConnectionError:

                st.error(
                    "Could not connect to the FastAPI backend. "
                    "Make sure the backend is running."
                )